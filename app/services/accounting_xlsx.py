from __future__ import annotations

from io import BytesIO
from typing import Any, Iterable
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app import models
from app.services.accounting_report import (
    FIXED_EXPENSE_KEYS,
    accounting_expense_entries,
    accounting_expense_payment_state,
    get_contract_expense_history,
    parse_number,
)

EXPENSE_LABELS = {
    "prosthetist_work": "Работа протезиста",
    "patient_travel": "Проезд пациента",
    "patient_accommodation": "Проживание пациента",
    "patient_meals": "Питание пациента",
    "patient_payment": "Выплата пациенту",
    "other_expenses": "Прочие расходы",
    "agency_expenses": "Агентские расходы",
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")


def _client_expense_history(client: models.Client, field_key: str, legacy_value: float) -> dict[str, Any]:
    entries = [dict(item) for item in accounting_expense_entries(client, field_key)]
    state = accounting_expense_payment_state(client, field_key) if field_key == "prosthetist_work" else "paid"
    if not entries and legacy_value > 0:
        entries = [{
            "id": f"legacy-{field_key}",
            "amount": float(legacy_value),
            "description": "Сумма до включения детализации",
            "created_at": None,
            "username": None,
            "paid": state == "paid",
        }]
    if field_key == "prosthetist_work":
        fallback_paid = state != "unpaid"
        for entry in entries:
            entry.setdefault("paid", fallback_paid)
    return {"entries": entries, "paid": state == "paid"}
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
THIN = Side(style="thin", color="D9E1F2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _style_sheet(ws, *, freeze: str = "A2") -> None:
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
    for column in range(1, ws.max_column + 1):
        values = [str(ws.cell(row=row, column=column).value or "") for row in range(1, min(ws.max_row, 100) + 1)]
        width = min(max(max((len(v) for v in values), default=8) + 2, 11), 38)
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.row_dimensions[1].height = 34


def _append_main_client_sheet(wb: Workbook, report: dict[str, Any]) -> None:
    ws = wb.active
    ws.title = "По клиентам"
    custom_fields = report.get("custom_fields", [])
    custom_by_key = {str(item.get("key")): item for item in custom_fields}
    headers = [
        "№", "Пациент", "Агент", "Дата пробития", "ТСР", "Сертификат",
        "Стоимость комплектующих", "Комплектующие со склада",
        *[EXPENSE_LABELS[key] for key in FIXED_EXPENSE_KEYS],
        *[str(field.get("label") or field.get("key")) for field in custom_fields],
        "НДС", "Налог", "Эквайринг", "Расходы", "Прибыль",
    ]
    ws.append(headers)
    for index, row in enumerate(report.get("rows", []), 1):
        client = row.get("client", {})
        amounts = row.get("amounts", {})
        expenses = row.get("expenses", {})
        custom_values = row.get("custom_values", {})
        certificates = row.get("certificates", [])
        cert_dates = ", ".join(sorted({str(item.get("date")) for item in certificates if item.get("date")}))
        tsr = ", ".join(dict.fromkeys(str(item.get("tsr") or "") for item in certificates if item.get("tsr")))
        fixed_total = sum(float(expenses.get(key) or 0) for key in FIXED_EXPENSE_KEYS)
        custom_total = sum(float(custom_values.get(str(field.get("key"))) or 0) for field in custom_fields if field.get("type") == "number")
        total_expenses = float(amounts.get("cost") or 0) + fixed_total + custom_total + float(amounts.get("vat") or 0) + float(amounts.get("tax") or 0) + float(amounts.get("acquiring") or 0)
        ws.append([
            index,
            client.get("full_name") or client.get("short_name") or "",
            client.get("agent_name") or client.get("agent") or "",
            cert_dates,
            tsr,
            float(amounts.get("revenue") or 0),
            float(amounts.get("cost") or 0),
            float(amounts.get("stock_reused_cost") or 0),
            *[float(expenses.get(key) or 0) for key in FIXED_EXPENSE_KEYS],
            *[custom_values.get(str(field.get("key")), "") for field in custom_fields],
            float(amounts.get("vat") or 0),
            float(amounts.get("tax") or 0),
            float(amounts.get("acquiring") or 0),
            total_expenses,
            float(amounts.get("profit") or 0),
        ])
    _style_sheet(ws)
    for col in range(6, ws.max_column + 1):
        for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
            for item in cell:
                if isinstance(item.value, (int, float)):
                    item.number_format = '#,##0.00 "₽"'


def _append_client_details(wb: Workbook, db: Session, report: dict[str, Any]) -> None:
    ws = wb.create_sheet("Расходы · по клиентам")
    ws.append(["Пациент", "Статья", "Сумма", "Описание", "Статус оплаты", "Дата", "Пользователь"])
    custom_fields = [field for field in report.get("custom_fields", []) if field.get("type") == "number"]
    keys = [(key, EXPENSE_LABELS[key]) for key in FIXED_EXPENSE_KEYS] + [
        (f"custom:{field.get('key')}", str(field.get("label") or field.get("key"))) for field in custom_fields
    ]
    for row in report.get("rows", []):
        client = row.get("client", {})
        client_id = client.get("client_id")
        if not client_id:
            continue
        client_model = db.get(models.Client, UUID(str(client_id)))
        if not client_model:
            continue
        expenses = row.get("expenses", {})
        custom_values = row.get("custom_values", {})
        for field_key, label in keys:
            legacy_value = (
                float(expenses.get(field_key) or 0)
                if field_key in FIXED_EXPENSE_KEYS
                else float(custom_values.get(field_key.split(":", 1)[1]) or 0)
            )
            history = _client_expense_history(client_model, field_key, legacy_value)
            for entry in history.get("entries", []):
                paid_label = ""
                if field_key == "prosthetist_work":
                    paid_label = "Оплачено" if bool(entry.get("paid", history.get("paid", True))) else "Не оплачено"
                ws.append([
                    client.get("full_name") or client.get("short_name") or "",
                    label,
                    float(entry.get("amount") or 0),
                    entry.get("description") or "",
                    paid_label,
                    entry.get("created_at") or "",
                    entry.get("username") or "",
                ])
    _style_sheet(ws)
    for cell in ws["C"][1:]:
        cell.number_format = '#,##0.00 "₽"'


def build_clients_accounting_xlsx(db: Session, report: dict[str, Any]) -> bytes:
    wb = Workbook()
    _append_main_client_sheet(wb, report)
    _append_client_details(wb, db, report)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def _append_contract_main_sheet(wb: Workbook, report: dict[str, Any]) -> None:
    ws = wb.active
    ws.title = "По договорам"
    custom_fields = report.get("custom_fields", [])
    headers = [
        "№", "Договор", "Пациент", "Дата договора", "Дата пробития", "Сертификат",
        "Стоимость комплектующих", "Комплектующие со склада",
        *[EXPENSE_LABELS[key] for key in FIXED_EXPENSE_KEYS],
        *[str(field.get("label") or field.get("key")) for field in custom_fields],
        "НДС", "Налог", "Эквайринг", "Расходы", "Прибыль",
    ]
    ws.append(headers)
    for index, row in enumerate(report.get("rows", []), 1):
        document = row.get("document", {})
        client = row.get("client", {})
        amounts = row.get("amounts", {})
        custom_values = row.get("custom_values", {})
        fixed_total = sum(float(amounts.get(key) or 0) for key in FIXED_EXPENSE_KEYS)
        custom_total = sum(float(custom_values.get(str(field.get("key"))) or 0) for field in custom_fields if field.get("type") == "number")
        total_expenses = float(amounts.get("modules_cost") or 0) + fixed_total + custom_total + float(amounts.get("vat") or 0) + float(amounts.get("tax") or 0) + float(amounts.get("acquiring") or 0)
        ws.append([
            index, document.get("filename") or document.get("document_number") or "", client.get("full_name") or "",
            document.get("date") or document.get("created_at") or "", document.get("certificate_date") or "",
            float(amounts.get("certificate") or 0), float(amounts.get("modules_cost") or 0), float(amounts.get("stock_reused_cost") or 0),
            *[float(amounts.get(key) or 0) for key in FIXED_EXPENSE_KEYS],
            *[custom_values.get(str(field.get("key")), "") for field in custom_fields],
            float(amounts.get("vat") or 0), float(amounts.get("tax") or 0), float(amounts.get("acquiring") or 0),
            total_expenses, float(amounts.get("profit") or 0),
        ])
    _style_sheet(ws)
    for row in ws.iter_rows(min_row=2, min_col=6):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.00 "₽"'


def _append_contract_details(wb: Workbook, db: Session, report: dict[str, Any]) -> None:
    ws = wb.create_sheet("Расходы · по договорам")
    ws.append(["Договор", "Пациент", "Статья", "Сумма", "Описание", "Статус оплаты", "Дата", "Пользователь"])
    custom_fields = [field for field in report.get("custom_fields", []) if field.get("type") == "number"]
    keys = [(key, EXPENSE_LABELS[key]) for key in FIXED_EXPENSE_KEYS] + [
        (f"custom:{field.get('key')}", str(field.get("label") or field.get("key"))) for field in custom_fields
    ]
    for row in report.get("rows", []):
        document = row.get("document", {})
        document_id = document.get("document_id")
        if not document_id:
            continue
        for field_key, label in keys:
            try:
                history = get_contract_expense_history(db, UUID(str(document_id)), field_key)
            except (ValueError, TypeError):
                continue
            for entry in history.get("entries", []):
                paid_label = ""
                if field_key == "prosthetist_work":
                    paid_label = "Оплачено" if bool(entry.get("paid", history.get("paid", True))) else "Не оплачено"
                ws.append([
                    document.get("filename") or document.get("document_number") or "",
                    row.get("client", {}).get("full_name") or "",
                    label, float(entry.get("amount") or 0), entry.get("description") or "", paid_label,
                    entry.get("created_at") or "", entry.get("username") or "",
                ])
    _style_sheet(ws)
    for cell in ws["D"][1:]:
        cell.number_format = '#,##0.00 "₽"'


def build_contracts_accounting_xlsx(db: Session, report: dict[str, Any]) -> bytes:
    wb = Workbook()
    _append_contract_main_sheet(wb, report)
    _append_contract_details(wb, db, report)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
