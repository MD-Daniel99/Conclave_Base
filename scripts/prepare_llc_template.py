"""Build the approved ООО «Смарт Движение» docxtpl template.

The source is kept next to the generated template so the markup can be
reproduced after future legal-text edits.
"""
from __future__ import annotations

from pathlib import Path
import re

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "app/templates/llc_contract_source_20260730.docx"
TARGET = ROOT / "app/templates/llc_contract.docx"


APPENDIX_RE = re.compile(r"^Приложение\s+([1-6])\s+от\b", re.IGNORECASE)


def _is_removable_empty_paragraph(element) -> bool:
    """Return True only for a plain empty spacer paragraph."""
    from docx.oxml.ns import qn

    if element.tag != qn("w:p"):
        return False
    protected = (
        ".//w:sectPr | .//w:drawing | .//w:pict | .//w:object | "
        ".//w:fldChar | .//w:instrText | .//w:br | .//w:bookmarkStart"
    )
    if element.xpath(protected):
        return False
    text = "".join(element.xpath(".//w:t/text()"))
    return not text.strip()


def _remove_spacers_before(paragraph) -> int:
    """Delete the legacy blank lines used to imitate page positioning."""
    removed = 0
    node = paragraph._p.getprevious()
    while node is not None and _is_removable_empty_paragraph(node):
        previous = node.getprevious()
        node.getparent().remove(node)
        removed += 1
        node = previous
    return removed


def enforce_appendix_page_starts(doc: Document) -> None:
    """Anchor applications 1-6 to real page boundaries.

    The approved source used many blank paragraphs to visually push each
    appendix to the next page. After docxtpl inserts longer customer data or
    expands TSR tables, those blank lines no longer correspond to a page and
    the next heading drifts upward or downward. We remove only safe, empty
    spacer paragraphs immediately before each appendix and set the semantic
    w:pageBreakBefore paragraph property instead.
    """
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    matches = []
    for paragraph in list(doc.paragraphs):
        text = " ".join(paragraph.text.split())
        match = APPENDIX_RE.match(text)
        if match:
            matches.append((int(match.group(1)), paragraph))

    numbers = [number for number, _ in matches]
    if numbers != [1, 2, 3, 4, 5, 6]:
        raise RuntimeError(
            "Ожидались приложения 1-6, найдены: "
            + ", ".join(str(number) for number in numbers)
        )

    for _, paragraph in matches:
        _remove_spacers_before(paragraph)
        paragraph.paragraph_format.page_break_before = True
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.widow_control = True

        # Keep the appendix number, contract reference and title together.
        node = paragraph._p
        kept = 0
        while kept < 3:
            node = node.getnext()
            if node is None or node.tag != qn("w:p"):
                break
            next_paragraph = Paragraph(node, paragraph._parent)
            next_paragraph.paragraph_format.keep_with_next = True
            next_paragraph.paragraph_format.widow_control = True
            kept += 1


def set_paragraph(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def replace_first(doc: Document, needle: str, replacement: str) -> None:
    for paragraph in doc.paragraphs:
        if needle in paragraph.text:
            set_paragraph(paragraph, replacement)
            return
    raise RuntimeError(f"Не найден абзац шаблона: {needle}")


def replace_all(doc: Document, needle: str, replacement: str) -> None:
    found = False
    for paragraph in doc.paragraphs:
        if needle in paragraph.text:
            set_paragraph(paragraph, replacement)
            found = True
    if not found:
        raise RuntimeError(f"Не найдены абзацы шаблона: {needle}")


def set_cell(cell, text: str) -> None:
    cell.text = text


def set_loop_table(table, values: list[str], *, total_values: list[str] | None = None) -> None:
    if len(table.rows) < 4:
        raise RuntimeError("В таблице недостаточно строк для цикла docxtpl")
    marker_row, data_row, end_row = table.rows[1:4]
    for cell in marker_row.cells:
        set_cell(cell, "")
    set_cell(marker_row.cells[0], "{%tr for item in contract_items %}")
    for cell, value in zip(data_row.cells, values):
        set_cell(cell, value)
    for cell in end_row.cells:
        set_cell(cell, "")
    set_cell(end_row.cells[0], "{%tr endfor %}")
    if total_values is not None:
        for cell, value in zip(table.rows[4].cells, total_values):
            set_cell(cell, value)


def replace_appendix_headers(doc: Document) -> None:
    appendix_number = 0
    for paragraph in doc.paragraphs:
        text = " ".join(paragraph.text.split())
        if text.startswith("Приложение ") and " от " in text:
            appendix_number += 1
            set_paragraph(paragraph, f"Приложение {appendix_number} от {{{{ДатаПриложения}}}}")
        elif text.startswith(("к Договору №", "К Договору №")):
            prefix = "К" if text.startswith("К ") else "к"
            set_paragraph(paragraph, f"{prefix} Договору № {{{{НомерДоговора}}}} от {{{{ДатаДоговора}}}}")


def main() -> None:
    doc = Document(SOURCE)

    # Header block is stored in the first table in the approved document.
    set_cell(
        doc.tables[0].rows[0].cells[1],
        "Договор № {{НомерДоговораБуквы}}/{{НомерДоговораЧисло}}\n"
        "на выполнение работ по изготовлению и поставку технических средств реабилитации "
        "от {{ДатаДоговора}} г.\nг. Москва",
    )

    replace_first(
        doc,
        "ООО «Смарт Движение», в лице Генерального директора",
        "ООО «Смарт Движение», в лице Генерального директора Абашеева Георгия Дмитриевича, "
        "действующего на основании Устава, именуемое в дальнейшем «Исполнитель», с одной стороны, "
        "и {{ФИОЗаказчика}}, именуемый(ая) в дальнейшем «Потребитель», с другой стороны, совместно "
        "именуемые «Стороны», заключили настоящий Договор о нижеследующем:",
    )
    replace_first(
        doc,
        "2.1. Стоимость выполнения работ и поставки ТСР",
        "2.1. Стоимость выполнения работ и поставки ТСР по настоящему Договору составляет "
        "{{Сумма}} ({{СуммаПрописью}}) рублей, НДС не облагается на основании п.1. ч.2 ст. 149 "
        "Налогового кодекса Российской Федерации.",
    )
    replace_all(
        doc,
        "_____________________ /Абашеев Г.Д./",
        "_____________________ /Абашеев Г.Д./\t\t______________________ /{{ИнициалыЗаказчикаКратко}}/",
    )
    replace_appendix_headers(doc)
    replace_first(
        doc,
        "ТСР изготавливаются на основании технических характеристик",
        "ТСР изготавливаются на основании технических характеристик, предусмотренных Индивидуальной "
        "программой реабилитации или абилитации (ИПРА) инвалида или Программой реабилитации "
        "пострадавшего в результате несчастного случая на производстве и профессионального заболевания "
        "(ПРП) № {{НомерИПРА}} от {{ДатаВыдачиИПРА}}.",
    )
    replace_first(doc, "Акт сдачи-приемки работ №", "Акт сдачи-приемки работ № {{НомерАкта}}")
    replace_first(doc, "«_____»________________20___ г.", "{{ДатаАкта}} г.")
    replace_first(
        doc,
        "ООО «Смарт Движение», в лице Генерального директора Абашеева Георгия Дмитриевича, действующего на основании Устава, именуемое в дальнейшем «Исполнитель», с одной стороны, и гр.",
        "ООО «Смарт Движение», в лице Генерального директора Абашеева Георгия Дмитриевича, "
        "действующего на основании Устава, именуемое в дальнейшем «Исполнитель», с одной стороны, и гр. "
        "{{ФИОЗаказчика}}, дата рождения {{ДатаРожденияЗаказчика}} г., паспорт: серия {{ПаспортСерия}} "
        "№ {{ПаспортНомер}}, выдан {{КемВыданПаспорт}}, {{ДатаВыдачиПаспорта}} г., код подразделения "
        "{{КодПодразделения}}, адрес места жительства: {{АдресРегистрацииЗаказчика}}, именуемый(ая) "
        "в дальнейшем «Потребитель», с другой стороны, совместно именуемые «Стороны», подписали "
        "настоящий Акт о нижеследующем:",
    )
    replace_first(
        doc,
        "Я, ___________________________________________________",
        "Я, {{ФИОЗаказчика}}, {{ДатаРожденияЗаказчика}} года рождения, место рождения "
        "{{МестоРожденияЗаказчика}}, паспорт: серия {{ПаспортСерия}} № {{ПаспортНомер}}, выдан "
        "{{КемВыданПаспорт}}, {{ДатаВыдачиПаспорта}} г., код подразделения {{КодПодразделения}}, "
        "зарегистрированный(ая) по адресу: {{АдресРегистрацииЗаказчика}}, даю своё согласие на "
        "обработку моих персональных данных.",
    )

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == "Пациент:":
            set_paragraph(paragraph, "Пациент: {{ФИОЗаказчика}}")
        elif text == "Дата рождения:":
            set_paragraph(paragraph, "Дата рождения: {{ДатаРожденияЗаказчика}}")
        elif text == "Адрес:":
            set_paragraph(paragraph, "Адрес: {{АдресРегистрацииЗаказчика}}")
        elif text.startswith("_______________ ________________________________________________________________"):
            set_paragraph(paragraph, "_______________ {{ФИОЗаказчика}}, {{Телефон_заказчика}}")

    requisites = doc.tables[1]
    set_cell(requisites.rows[0].cells[1], "Потребитель: {{ФИОЗаказчика}}")
    set_cell(requisites.rows[1].cells[1], "Паспорт: {{ПаспортПолностью}}")
    set_cell(requisites.rows[2].cells[1], "Адрес регистрации: {{АдресРегистрацииЗаказчика}}")
    set_cell(requisites.rows[3].cells[1], "СНИЛС: {{СНИЛСЗаказчика}}")
    set_cell(requisites.rows[4].cells[1], "Телефон: {{Телефон_заказчика}}")
    set_cell(requisites.rows[5].cells[1], "Дата выдачи ИПРА: {{ДатаВыдачиИПРА}}")
    set_cell(requisites.rows[6].cells[1], "Номер ИПРА: {{НомерИПРА}}")

    set_loop_table(doc.tables[2], ["{{item.number}}", "{{item.description}}", "{{item.quantity}}"])
    set_loop_table(
        doc.tables[3],
        ["{{item.name_and_code}}", "{{item.quantity}}", "{{item.price}}", "{{item.total_price}}"],
        total_values=["ИТОГО:", "{{ИтогоКоличество}}", "", "{{ИтогоСтоимость}}"],
    )
    set_loop_table(
        doc.tables[4],
        ["{{item.number}}", "{{item.description}}", "{{item.quantity}}", "{{item.components_text}}"],
    )

    enforce_appendix_page_starts(doc)

    TARGET.parent.mkdir(parents=True, exist_ok=True)
    doc.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    main()
